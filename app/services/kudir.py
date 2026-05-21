from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.models import BankTransaction, TransactionDirection


MONTH_NAMES_RU = {
    1: "Январь",
    2: "Февраль",
    3: "Март",
    4: "Апрель",
    5: "Май",
    6: "Июнь",
    7: "Июль",
    8: "Август",
    9: "Сентябрь",
    10: "Октябрь",
    11: "Ноябрь",
    12: "Декабрь",
}


@dataclass(slots=True)
class KudirRow:
    row_number: int
    operation_date: date
    document_number: str
    income: Decimal | None
    expense: Decimal | None
    basis: str


def _classify_basis(tx: BankTransaction) -> str:
    ref_lower = (tx.reference or "").lower()
    if any(kw in ref_lower for kw in ("комисс", "комміс", "ркo", "рко", "обслужив")):
        return settings.kudir_expense_bank_fee_basis
    if tx.direction == TransactionDirection.outgoing:
        return settings.kudir_expense_third_party_basis
    return settings.kudir_income_basis


def _make_document_number(op_date: date, seq: int) -> str:
    return f"{op_date.strftime('%d.%m')}-{seq:02d}"


def build_kudir_rows(
    session: Session,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> list[KudirRow]:
    filters = []
    if date_from is not None:
        filters.append(BankTransaction.booked_at >= datetime(date_from.year, date_from.month, date_from.day))
    if date_to is not None:
        next_day = datetime(date_to.year, date_to.month, date_to.day, 23, 59, 59)
        filters.append(BankTransaction.booked_at <= next_day)

    transactions = session.scalars(
        select(BankTransaction)
        .where(*filters)
        .order_by(BankTransaction.booked_at, BankTransaction.id)
    ).all()

    day_counters: dict[date, int] = defaultdict(int)
    rows: list[KudirRow] = []

    for idx, tx in enumerate(transactions, start=1):
        op_date = tx.booked_at.date()
        day_counters[op_date] += 1
        seq = day_counters[op_date]

        basis = _classify_basis(tx)
        is_income = tx.direction == TransactionDirection.incoming and basis == settings.kudir_income_basis

        rows.append(
            KudirRow(
                row_number=idx,
                operation_date=op_date,
                document_number=_make_document_number(op_date, seq),
                income=tx.amount if is_income else None,
                expense=tx.amount if not is_income else None,
                basis=basis,
            )
        )

    return rows


def _apply_header(ws, ip_name: str, inn: str, tax_year: int, tax_system: str) -> None:
    header_font = Font(bold=True, size=12)
    info_font = Font(size=10)

    ws.merge_cells("A1:E1")
    ws["A1"] = "Книга учёта доходов и расходов"
    ws["A1"].font = header_font

    ws.merge_cells("A2:E2")
    ws["A2"] = ip_name
    ws["A2"].font = info_font

    ws.merge_cells("A3:E3")
    ws["A3"] = f"ИНН {inn}"
    ws["A3"].font = info_font

    ws.merge_cells("A4:E4")
    ws["A4"] = f"Налоговый период: {tax_year} год"
    ws["A4"].font = info_font

    ws.merge_cells("A5:E5")
    ws["A5"] = f"Система налогообложения: {tax_system}"
    ws["A5"].font = info_font


def _apply_table_header(ws, start_row: int = 1) -> None:
    headers = {
        "F": "№",
        "G": "Дата операции",
        "H": "№ документа",
        "I": "Доход, ₽",
        "J": "Расход, ₽",
        "K": "Основание",
    }
    header_font = Font(bold=True, size=10)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_letter, title in headers.items():
        cell = ws[f"{col_letter}{start_row}"]
        cell.value = title
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_data_rows(ws, rows: list[KudirRow], start_row: int = 2) -> None:
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    date_font = Font(size=10)
    number_format_rub = '#,##0.00'

    for i, row in enumerate(rows):
        excel_row = start_row + i

        cell_num = ws[f"F{excel_row}"]
        cell_num.value = row.row_number
        cell_num.border = thin_border
        cell_num.alignment = Alignment(horizontal="center")

        cell_date = ws[f"G{excel_row}"]
        cell_date.value = row.operation_date.strftime("%d.%m.%Y")
        cell_date.font = date_font
        cell_date.border = thin_border
        cell_date.alignment = Alignment(horizontal="center")

        cell_doc = ws[f"H{excel_row}"]
        cell_doc.value = row.document_number
        cell_doc.border = thin_border
        cell_doc.alignment = Alignment(horizontal="center")

        cell_income = ws[f"I{excel_row}"]
        cell_income.value = float(row.income) if row.income is not None else "–"
        if row.income is not None:
            cell_income.number_format = number_format_rub
        cell_income.border = thin_border
        cell_income.alignment = Alignment(horizontal="right")

        cell_expense = ws[f"J{excel_row}"]
        cell_expense.value = float(row.expense) if row.expense is not None else "–"
        if row.expense is not None:
            cell_expense.number_format = number_format_rub
        cell_expense.border = thin_border
        cell_expense.alignment = Alignment(horizontal="right")

        cell_basis = ws[f"K{excel_row}"]
        cell_basis.value = row.basis
        cell_basis.border = thin_border
        cell_basis.alignment = Alignment(wrap_text=True)


def _set_column_widths(ws) -> None:
    widths = {"F": 6, "G": 16, "H": 14, "I": 16, "J": 16, "K": 60}
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
    for col_letter in ("A", "B", "C", "D", "E"):
        ws.column_dimensions[col_letter].width = 12


def generate_kudir_xlsx(
    rows: list[KudirRow],
    *,
    ip_name: str | None = None,
    inn: str | None = None,
    tax_year: int | None = None,
    tax_system: str | None = None,
    month: int | None = None,
) -> BytesIO:
    wb = Workbook()
    ws = wb.active

    effective_ip_name = ip_name or settings.kudir_ip_name
    effective_inn = inn or settings.kudir_inn
    effective_tax_year = tax_year or settings.kudir_tax_year or datetime.now().year
    effective_tax_system = tax_system or settings.kudir_tax_system

    if month and 1 <= month <= 12:
        ws.title = MONTH_NAMES_RU[month]
    else:
        ws.title = "КУДиР"

    _apply_header(ws, effective_ip_name, effective_inn, effective_tax_year, effective_tax_system)
    _apply_table_header(ws, start_row=1)
    _write_data_rows(ws, rows, start_row=2)
    _set_column_widths(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_kudir_xlsx_for_period(
    session: Session,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> tuple[BytesIO, str]:
    rows = build_kudir_rows(session, date_from=date_from, date_to=date_to)

    month: int | None = None
    if date_from and date_to and date_from.month == date_to.month and date_from.year == date_to.year:
        month = date_from.month

    buf = generate_kudir_xlsx(rows, month=month)

    if date_from and date_to:
        filename = f"kudir_{date_from.isoformat()}_{date_to.isoformat()}.xlsx"
    else:
        filename = "kudir.xlsx"

    return buf, filename


def save_kudir_xlsx(
    session: Session,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> Path:
    buf, filename = generate_kudir_xlsx_for_period(session, date_from=date_from, date_to=date_to)
    artifacts_dir = Path(settings.kudir_artifacts_dir)
    artifacts_dir.mkdir(parents=True, exist_ok=True)
    output_path = artifacts_dir / filename
    output_path.write_bytes(buf.getvalue())
    return output_path
